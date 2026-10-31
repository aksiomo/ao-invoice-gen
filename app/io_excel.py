"""把 xlsx / OLE2 xls / WPS 导出的 HTML 伪 xls 统一读成二维字符串矩阵。

只做确定性的表格读取，不做任何业务列识别（业务识别交给 Agent A）。
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

from bs4 import BeautifulSoup

HEADER_KEYWORDS = ("存货名称", "商品名称", "货物名称", "名称", "数量", "单价", "金额")


def detect_format(path: str | Path) -> str:
    with open(path, "rb") as f:
        head = f.read(2048)
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    text = head.lstrip().lower()
    if text.startswith(b"<html") or b"<table" in head.lower() or b"<meta" in text:
        return "html"
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return "xlsx"
    if suffix in (".xls", ".xlt"):
        return "xls"
    return "html"


def read_matrix(path: str | Path) -> list[list[str]]:
    fmt = detect_format(path)
    if fmt == "xlsx":
        return _read_xlsx(path)
    if fmt == "xls":
        return _read_xls(path)
    return _read_html(path)


def _fmt_scalar(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g").rstrip("0").rstrip(".")
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.isoformat()
    return str(value)


def _score_sheet(rows: list[list[str]]) -> int:
    score = 0
    nonempty = 0
    for row in rows:
        for cell in row:
            cell = cell.strip()
            if cell:
                nonempty += 1
                if any(k in cell for k in HEADER_KEYWORDS):
                    score += 3
    return score + nonempty // 100


def _read_xlsx(path: str | Path) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    best: list[list[str]] = []
    best_score = -1
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_fmt_scalar(v).strip() for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            rows.append(cells)
        score = _score_sheet(rows)
        if score > best_score:
            best_score, best = score, rows
    wb.close()
    return best


def _read_xls(path: str | Path) -> list[list[str]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    best: list[list[str]] = []
    best_score = -1
    for sheet in book.sheets():
        rows: list[list[str]] = []
        for r in range(sheet.nrows):
            cells = []
            for c in range(sheet.ncols):
                value = sheet.cell_value(r, c)
                if isinstance(value, float) and value.is_integer():
                    cells.append(str(int(value)))
                else:
                    cells.append(_fmt_scalar(value).strip())
            while cells and cells[-1] == "":
                cells.pop()
            rows.append(cells)
        score = _score_sheet(rows)
        if score > best_score:
            best_score, best = score, rows
    return best


def _read_html(path: str | Path) -> list[list[str]]:
    html = Path(path).read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "lxml")
    matrix: list[list[str]] = []
    for table in soup.find_all("table"):
        matrix.extend(_expand_html_table(table))
    if not matrix:
        soup = BeautifulSoup(html, "html.parser")
        for table in soup.find_all("table"):
            matrix.extend(_expand_html_table(table))
    return matrix


def _cell_text(cell) -> str:
    text = cell.get_text(separator=" ", strip=True)
    return text.replace("\xa0", " ").replace("\u3000", " ").strip()


def _expand_html_table(table) -> list[list[str]]:
    rows_out: list[list[str]] = []
    carry: dict[int, list] = {}
    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"])
        out: list[str] = []
        col = 0
        for cell in cells:
            while col in carry and carry[col][1] > 0:
                out.append(carry[col][0])
                carry[col][1] -= 1
                col += 1
            text = _cell_text(cell)
            try:
                rowspan = max(1, int(cell.get("rowspan", 1)))
            except (TypeError, ValueError):
                rowspan = 1
            try:
                colspan = max(1, int(cell.get("colspan", 1)))
            except (TypeError, ValueError):
                colspan = 1
            for k in range(colspan):
                out.append(text)
                if rowspan > 1:
                    carry[col] = [text, rowspan - 1]
                col += 1
        max_col = max([col, *[c + 1 for c, v in carry.items() if v[1] > 0]], default=0)
        while col < max_col:
            if col in carry and carry[col][1] > 0:
                out.append(carry[col][0])
                carry[col][1] -= 1
            else:
                out.append("")
            col += 1
        while out and out[-1] == "":
            out.pop()
        rows_out.append(out)
    return rows_out
