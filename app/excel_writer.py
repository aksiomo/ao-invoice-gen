"""基于官方《发票开具项目信息导入模板.xlsx》生成结果文件。

数据写入 sheet `1-明细模板` 第 4 行起，其余 sheet 与说明/表头行原样保留。
列序：项目名称 | 商品和服务税收分类编码 | 规格型号 | 单位 | 商品数量 |
      商品单价 | 金额 | 税率 | 折扣金额 | 优惠政策类型 | 煤炭种类
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.config import settings
from app.schemas import FinalRow, PurchaseItem

DATA_SHEET = "1-明细模板"
START_ROW = 4


def _num(value: Decimal | None, places: int | None = None):
    if value is None:
        return None
    if places is not None:
        quant = Decimal(1).scaleb(-places)
        return float(value.quantize(quant))
    return float(value)


def build_final_rows(
    items: list[PurchaseItem],
    classified: dict[str, object],
    *,
    tolerance: float | None = None,
) -> tuple[list[FinalRow], list[str]]:
    tol = Decimal(str(settings.amount_tolerance if tolerance is None else tolerance))
    rows: list[FinalRow] = []
    warnings: list[str] = []
    for item in items:
        c = classified.get(item.name)
        if c is None:
            raise KeyError(f"商品缺少分类结果：{item.name}")
        amount = item.amount
        amount_warning = False
        if item.quantity is not None and item.unit_price is not None:
            computed = item.quantity * item.unit_price
            if amount is None:
                amount = computed.quantize(Decimal("0.01"))
            elif abs(amount - computed.quantize(Decimal("0.01"))) > tol:
                amount_warning = True
                warnings.append(
                    f"“{item.name}”金额 {amount} 与 数量×单价 "
                    f"{computed.quantize(Decimal('0.01'))} 不一致"
                )
        rows.append(
            FinalRow(
                name=item.name,
                code=c.code,
                category_name=c.category_name,
                unit=item.unit,
                quantity=item.quantity,
                unit_price=item.unit_price,
                amount=amount,
                tax_rate=Decimal(str(settings.fixed_tax_rate)),
                from_cache=c.from_cache,
                receipt_no=item.receipt_no,
                amount_warning=amount_warning,
            )
        )
    return rows, warnings


def write_workbook(rows: list[FinalRow], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(settings.template_path)
    ws = wb[DATA_SHEET]

    for i, row in enumerate(rows):
        r = START_ROW + i
        ws.cell(r, 1, row.category_name)
        ws.cell(r, 2, row.code)
        ws.cell(r, 3, row.name)
        ws.cell(r, 4, row.unit)
        ws.cell(r, 5, _num(row.quantity))
        ws.cell(r, 6, _num(row.unit_price))
        ws.cell(r, 7, _num(row.amount, 2))
        ws.cell(r, 8, float(row.tax_rate))

    wb.save(output_path)
    return output_path
