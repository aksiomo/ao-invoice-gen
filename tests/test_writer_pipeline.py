from __future__ import annotations

from decimal import Decimal

import openpyxl

from app.excel_writer import build_final_rows, write_workbook
from app.io_excel import read_matrix
from app.schemas import ClassifiedItem, PurchaseItem

from .conftest import PURCHASE_FILE


def _item(name, unit="个", qty="1", price="10", amount="10"):
    return PurchaseItem(
        name=name,
        unit=unit,
        quantity=Decimal(qty),
        unit_price=Decimal(price),
        amount=Decimal(amount),
    )


def test_writer_preserves_template_and_writes_rows(tmp_path):
    items = [_item("壁扇", "台", "1", "138", "138"), _item("600*600天花灯", "盏", "10", "60", "600")]
    classified = {
        "壁扇": ClassifiedItem(name="壁扇", code="1090416010000000000", category_name="家用电风扇"),
        "600*600天花灯": ClassifiedItem(
            name="600*600天花灯", code="1090424010000000000", category_name="室内照明灯具"
        ),
    }
    rows, warnings = build_final_rows(items, classified)
    out = write_workbook(rows, tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["1-明细模板", "excelVersion", "xzqhdm", "2-特定业务信息"]
    ws = wb["1-明细模板"]
    assert ws.cell(3, 1).value == "项目名称"
    assert [ws.cell(4, c).value for c in range(1, 9)] == [
        "家用电风扇",
        "1090416010000000000",
        "壁扇",
        "台",
        1.0,
        138.0,
        138.0,
        0.13,
    ]
    assert ws.cell(5, 8).value == 0.13
    assert not warnings


def test_amount_mismatch_warning():
    items = [_item("壁扇", qty="2", price="10", amount="25")]
    classified = {
        "壁扇": ClassifiedItem(name="壁扇", code="1090416010000000000", category_name="家用电风扇")
    }
    rows, warnings = build_final_rows(items, classified)
    assert warnings and rows[0].amount_warning


def test_name_invariant_across_pipeline(tmp_path, monkeypatch):
    """核心不变量：不调用 LLM 时，输出规格型号列必须与输入商品名集合逐字符一致。"""
    from app import pipeline as pl
    from app.schemas import ClassifiedItem

    matrix = read_matrix(PURCHASE_FILE)

    fake_items = []
    import re

    for r in matrix:
        if len(r) >= 7 and re.fullmatch(r"\d+", r[0] or "") and r[1]:
            fake_items.append(
                PurchaseItem(
                    receipt_no="",
                    item_code=r[0],
                    name=r[1],
                    spec_input=r[2],
                    unit=r[3],
                    quantity=Decimal(r[4]) if r[4] else None,
                    unit_price=Decimal(r[5].replace(",", "")) if r[5] else None,
                    amount=Decimal(r[6].replace(",", "")) if r[6] else None,
                )
            )

    monkeypatch.setattr(pl, "read_matrix", lambda p: matrix)
    monkeypatch.setattr(pl, "extract_items", lambda m: fake_items)
    monkeypatch.setattr(
        pl,
        "classify_names",
        lambda names: (
            [
                ClassifiedItem(name=n, code="1090416010000000000", category_name="家用电风扇")
                for n in names
            ],
            0,
        ),
    )

    class _Job:
        def __init__(self):
            self.id = "testjob"
            self.filename = "采购入库单.XLS"
            self.stage = "queued"
            self.message = ""
            self.total_items = 0
            self.unique_names = 0
            self.cache_hits = 0
            self.warnings = []
            self.rows = []
            self.items = []
            self.output_path = tmp_path / "out.xlsx"
            self.error = ""

    job = _Job()
    pl.run_pipeline(job, PURCHASE_FILE)
    assert job.stage == "done", job.error

    wb = openpyxl.load_workbook(job.output_path)
    ws = wb["1-明细模板"]
    out_names = [ws.cell(r, 3).value for r in range(4, 4 + len(fake_items))]
    assert out_names == [it.name for it in fake_items]

    input_names = [it.name for it in fake_items]
    assert sorted(out_names) == sorted(input_names)
    assert ws.cell(4, 8).value == 0.13
